"""One project as a three-tab spreadsheet.

This is an export of measurements, not a model: every cell is a number that
was computed from the depth maps and the cluster file, so there are no
formulas and nothing to recalculate. What makes it trustworthy instead is the
first tab, which records every setting that produced the other two — without
it, two workbooks of the same project are not comparable and nobody can tell
why.

The numbers come from ``features.compute``, the same function the
cross-project page reads, so a figure here and a figure there cannot drift.
The browser's own copies of this arithmetic are pinned to it by a test.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import asdict
from pathlib import Path
from typing import Callable, List, Optional, Sequence, Tuple

from ..collect import bundle as B
from . import corrections as C
from . import features as F

FONT = 'Arial'
HEADER_FILL = 'FF1F2733'
BLOCK_FILL = 'FFE8ECF1'


class WorkbookError(Exception):
    """openpyxl is absent, or there is nothing to write."""


def _require_openpyxl():
    try:
        import openpyxl                     # noqa: F401
        from openpyxl.styles import Alignment, Font, PatternFill   # noqa: F401
    except ImportError:
        raise WorkbookError(
            'openpyxl is not installed. pip install "openpyxl>=3.1"')
    return openpyxl


# --------------------------------------------------------------- tab 2 and 3
# Column order is part of the contract: someone will build a script against
# these headings, so new columns go on the end rather than in the middle.

def _get(row: dict, *path, default=None):
    """Reach into nested metrics without a pile of .get() calls."""
    value = row
    for key in path:
        if not isinstance(value, dict):
            return default
        value = value.get(key)
        if value is None:
            return default
    return value


def _shape_columns(prefix: str, *path) -> List[Tuple[str, Callable]]:
    """The five numbers an ellipse is worth reporting, for one group."""
    return [
        (prefix + ' n', lambda r, p=path: _get(r, *p, 'n')),
        (prefix + ' area cm2', lambda r, p=path: _get(r, *p, 'areaCm2')),
        (prefix + ' aspect', lambda r, p=path: _get(r, *p, 'aspect')),
        (prefix + ' aspect if circular',
         lambda r, p=path: _get(r, *p, 'nullAspect')),
        (prefix + ' angle deg', lambda r, p=path: _get(r, *p, 'angleDeg')),
    ]


def trial_columns() -> List[Tuple[str, Callable]]:
    columns: List[Tuple[str, Callable]] = [
        ('projectID', lambda r: r['_projectID']),
        ('analysisID', lambda r: r['_analysisID']),
        ('tankID', lambda r: r['_tankID']),
        ('category', lambda r: r['_category']),
        ('trial', lambda r: r.get('trial')),
        ('days', lambda r: r.get('days')),
        ('first date', lambda r: r.get('from')),
        ('last date', lambda r: r.get('to')),

        ('castle volume cm3', lambda r: r.get('castleVolume')),
        ('pit volume cm3', lambda r: r.get('pitVolume')),
        ('total volume cm3', lambda r: r.get('totalVolume')),
        ('bower index', lambda r: r.get('bowerIndex')),
        ('castle area cm2', lambda r: r.get('castleArea')),
        ('pit area cm2', lambda r: r.get('pitArea')),
        ('volume per day cm3', lambda r: r.get('volumePerDay')),
        ('days to 90% of final', lambda r: r.get('daysTo90')),

        ('gross volume moved cm3', lambda r: r.get('grossVolume')),
        ('night volume cm3', lambda r: r.get('nightVolume')),
        ('day volume cm3', lambda r: r.get('dayVolume')),
        ('night share', lambda r: r.get('nightShare')),
        ('night rate / day rate', lambda r: r.get('nightRateRatio')),
        ('mean dark hours', lambda r: r.get('darkHoursMean')),
        ('mean light hours', lambda r: r.get('lightHoursMean')),

        ('split day index', lambda r: _get(r, 'split', 'dayIndex')),
        ('split date', lambda r: _get(r, 'split', 'date')),
        ('split gross fraction', lambda r: _get(r, 'split', 'fraction')),
    ]
    for label, key in (('first half', 'firstHalf'), ('second half', 'secondHalf')):
        columns += [
            (label + ' from', lambda r, k=key: _get(r, k, 'fromDate')),
            (label + ' to', lambda r, k=key: _get(r, k, 'toDate')),
            (label + ' days', lambda r, k=key: _get(r, k, 'days')),
            (label + ' gross share', lambda r, k=key: _get(r, k, 'grossShare')),
            (label + ' net volume cm3', lambda r, k=key: _get(r, k, 'totalVolume')),
            (label + ' castle cm3', lambda r, k=key: _get(r, k, 'castleVolume')),
            (label + ' pit cm3', lambda r, k=key: _get(r, k, 'pitVolume')),
            (label + ' index', lambda r, k=key: _get(r, k, 'bowerIndex')),
        ]

    columns.append(('events', lambda r: r.get('events')))
    columns.append(('events per day', lambda r: r.get('eventsPerDay')))
    for name, _bid in F.BEHAVIOUR:
        columns.append((name, lambda r, n=name: r.get(n)))
    columns.append(('scoop / spit', lambda r: r.get('scoopSpitRatio')))
    columns.append(('events placed in depth coords',
                    lambda r: r.get('placed')))

    columns += _shape_columns('scoop', 'scoopShape')
    columns += _shape_columns('spit', 'spitShape')
    columns += [
        ('spit area / scoop area', lambda r: r.get('spreadRatio')),
        ('centroids apart cm', lambda r: r.get('centroidGap')),
    ]
    columns += _shape_columns('first half scoop', 'firstHalfShapes', 'scoop')
    columns += _shape_columns('first half spit', 'firstHalfShapes', 'spit')
    columns += _shape_columns('second half scoop', 'secondHalfShapes', 'scoop')
    columns += _shape_columns('second half spit', 'secondHalfShapes', 'spit')

    columns += [
        ('spawns measured', lambda r: r.get('spawnsMeasured')),
        ('spawn depth mean cm', lambda r: r.get('spawnDepthMean')),
        ('spawn depth median cm', lambda r: r.get('spawnDepthMedian')),
        ('spawn depth q1 cm',
         lambda r: (r.get('spawnDepthQuartiles') or [None])[0]),
        ('spawn depth q3 cm',
         lambda r: (r.get('spawnDepthQuartiles') or [None, None, None])[2]),
        ('spawns over added sand %', lambda r: r.get('spawnOverCastle')),
        ('days marked new bower', lambda r: r.get('newBowerDays')),
        ('days marked wall building', lambda r: r.get('wallBuildingDays')),
    ]
    return columns


def day_columns() -> List[Tuple[str, Callable]]:
    columns: List[Tuple[str, Callable]] = [
        ('projectID', lambda r: r['_projectID']),
        ('analysisID', lambda r: r['_analysisID']),
        ('category', lambda r: r['_category']),
        ('trial', lambda r: r.get('trial')),
        ('day index', lambda r: r.get('index')),
        ('date', lambda r: r.get('date')),
        ('window from', lambda r: r.get('from')),
        ('window to', lambda r: r.get('to')),
        ('light hours', lambda r: r.get('lightHours')),
        ('dark hours before', lambda r: r.get('darkHours')),

        ('day castle cm3', lambda r: _get(r, 'light', 'castleVolume')),
        ('day pit cm3', lambda r: _get(r, 'light', 'pitVolume')),
        ('day total cm3', lambda r: _get(r, 'light', 'totalVolume')),
        ('day index value', lambda r: _get(r, 'light', 'bowerIndex')),
        ('night castle cm3', lambda r: _get(r, 'dark', 'castleVolume')),
        ('night pit cm3', lambda r: _get(r, 'dark', 'pitVolume')),
        ('night total cm3', lambda r: _get(r, 'dark', 'totalVolume')),
        ('night index value', lambda r: _get(r, 'dark', 'bowerIndex')),
        ('24 hour total cm3', lambda r: r.get('_total24')),
        ('night share', lambda r: r.get('_nightShare')),

        ('cumulative castle cm3', lambda r: _get(r, 'cumulative', 'castleVolume')),
        ('cumulative pit cm3', lambda r: _get(r, 'cumulative', 'pitVolume')),
        ('cumulative total cm3', lambda r: _get(r, 'cumulative', 'totalVolume')),
        ('cumulative index', lambda r: _get(r, 'cumulative', 'bowerIndex')),

        ('events', lambda r: _get(r, '_events', 'events')),
    ]
    for name, _bid in F.BEHAVIOUR:
        columns.append((name, lambda r, n=name: _get(r, '_events', n)))
    columns.append(('scoop / spit', lambda r: _get(r, '_events', 'scoopSpitRatio')))
    columns += _shape_columns('scoop', '_events', 'scoopShape')
    columns += _shape_columns('spit', '_events', 'spitShape')
    columns += [
        ('centroids apart cm', lambda r: _get(r, '_events', 'centroidGap')),
        ('spawns measured', lambda r: _get(r, '_events', 'spawnsMeasured')),
        ('spawn depth mean cm', lambda r: _get(r, '_events', 'spawnDepthMean')),
        ('spawn depth min cm', lambda r: _get(r, '_events', 'spawnDepthMin')),
        ('spawn depth max cm', lambda r: _get(r, '_events', 'spawnDepthMax')),
        ('marked new bower', lambda r: r.get('_newBower')),
        ('marked wall building', lambda r: r.get('_wallBuilding')),
    ]
    return columns


# -------------------------------------------------------------------- tab 1

def settings_rows(project_paths, metrics: dict, prep, manifest: dict,
                  category: str, source: str) -> List[Sequence]:
    """Parameter, value, units, where it came from.

    The fourth column is the point of the sheet: a reader a year from now has
    to be able to tell a setting someone chose from a constant of the rig, and
    a saved correction from a slider position at the moment of download.
    """
    params = metrics.get('params') or asdict(F.DEFAULTS)
    clusters = manifest.get('clusters') or {}
    rows: List[Sequence] = []

    def block(title):
        rows.append([])
        rows.append([title])

    def row(name, value, units='', origin=''):
        rows.append([name, value, units, origin])

    rows.append(['Parameter', 'Value', 'Units', 'Where it came from'])

    block('Analysis settings — these change the numbers')
    row('Confidence threshold', params['confidence'], 'probability', source)
    row('Hour window start', params['hour_from'], 'hour of day', source)
    row('Hour window end', params['hour_to'], 'hour of day', source)
    row('Sand moved threshold', params['threshold'], 'cm', source)
    row('Ellipse coverage', params['coverage'], 'share of events', source)
    row('Volume split point', params['split_at'], 'share of gross volume', source)
    row('Minimum events for a reported shape', params['min_for_shape'],
        'events', source)
    row('Residual k', prep.residual_k, 'MADs above the log median',
        'saved in prep.json')

    block('Display only — these change a picture, not a number')
    row('Event spread', params['spread'], 'depth px', source)

    block('Registration and crops')
    transform = prep.transform
    if transform:
        for r_index, line in enumerate(transform):
            row('Registration matrix row %d' % (r_index + 1),
                ', '.join('%.6g' % v for v in line), 'video px to depth px',
                'saved in prep.json')
    else:
        row('Registration matrix', 'none saved', '', 'saved in prep.json')
    for label, points in (('Depth crop', prep.depth_crop),
                          ('Video crop', prep.video_crop)):
        if points:
            for p_index, point in enumerate(points):
                row('%s corner %d' % (label, p_index + 1),
                    '%g, %g' % (point[0], point[1]), 'x, y px',
                    'saved in prep.json')
        else:
            row(label, 'none saved', '', 'saved in prep.json')

    block('Per trial')
    rows.append(['Trial', 'Excluded', 'Reason', 'Own registration',
                 'Own crops', 'Start adjust min', 'Stop adjust min'])
    for trial in manifest.get('trials', []):
        number = trial.get('number')
        override = prep.override(number)
        times = (prep.trials or {}).get(str(number))
        rows.append([
            number,
            'yes' if prep.is_excluded(number) else '',
            override.reason or '',
            'yes' if override.transform else '',
            'yes' if (override.depth_crop or override.video_crop) else '',
            times.start if times else 0,
            times.stop if times else 0,
        ])

    block('Measurement constants — what the rig is')
    row('Pixel length', F.PIXEL_CM, 'cm per depth px', 'fixed')
    row('Depth frame size',
        ' x '.join(str(v) for v in (manifest.get('depthSize') or [])) or '?',
        'px', 'from the bundle')
    row('Spawn depth histogram span', metrics.get('depthSpan'), 'cm either way',
        'fixed')
    row('Spawn depth histogram bins', metrics.get('depthBins'), 'bins', 'fixed')

    block('Provenance — whether two workbooks are comparable')
    row('projectID', manifest.get('projectID'), '', 'manifest')
    row('analysisID', manifest.get('analysisID'), '', 'manifest')
    row('tankID', manifest.get('tankID'), '', 'manifest')
    row('Category', category, '', 'analysis states file')
    row('Prep last saved by', prep.who, '', 'prep.json')
    row('Prep last saved at', prep.updated, '', 'prep.json')
    row('Prep schema', prep.schema, '', 'prep.json')
    row('Manifest schema', manifest.get('schema'), '', 'manifest')
    row('Collected at', manifest.get('built') or manifest.get('collected'),
        '', 'manifest')
    row('Collected from branch', manifest.get('branch'), '', 'manifest')
    row('Frames extracted', manifest.get('extracted'), 'frames', 'manifest')
    row('Frames missing from the archive', len(manifest.get('missing') or []),
        'frames', 'manifest')
    row('Values clipped at float16 max', manifest.get('clipped'), 'pixels',
        'manifest')
    row('Bundle size', manifest.get('bundleBytes'), 'bytes', 'manifest')
    row('Source archive size', manifest.get('sourceBytes'), 'bytes', 'manifest')
    row('Cluster detections', clusters.get('n'), 'rows', 'cluster file')
    row('Unclustered (LID -1)', clusters.get('unclustered'), 'rows',
        'cluster file')
    row('Clustered fraction', clusters.get('clusteredFraction'), '',
        'cluster file')
    row('Detections with no prediction', clusters.get('noPrediction'), 'rows',
        'cluster file')
    row('Detections with no clip', clusters.get('noClip'), 'rows',
        'cluster file')
    row('Metrics schema', metrics.get('schema'), '', 'this export')
    row('Workbook written', str(dt.datetime.now().replace(microsecond=0)),
        '', 'this export')

    rows.append([])
    rows.append(['Every value in the Trials and Days sheets was measured with '
                 'the settings above. There are no formulas in this workbook: '
                 'changing a cell changes nothing else.'])
    return rows


# -------------------------------------------------------------------- writing

def _style_header(sheet, row_index, width, openpyxl):
    from openpyxl.styles import Alignment, Font, PatternFill
    for column in range(1, width + 1):
        cell = sheet.cell(row=row_index, column=column)
        cell.font = Font(name=FONT, bold=True, color='FFFFFFFF')
        cell.fill = PatternFill('solid', fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical='top', wrap_text=True)


def _autofit(sheet, widths):
    from openpyxl.utils import get_column_letter
    for index, width in widths.items():
        sheet.column_dimensions[get_column_letter(index)].width = \
            min(42, max(9, width + 2))


def write_workbook(path: Path, project_paths, metrics: dict, category: str,
                   source: str = 'chosen at download') -> Path:
    """Write the three sheets. Returns the path written."""
    openpyxl = _require_openpyxl()
    from openpyxl.styles import Font, PatternFill

    manifest = B.read_manifest(project_paths.manifest) or {}
    prep = C.load(project_paths)

    book = openpyxl.Workbook()

    # --- tab 1 ------------------------------------------------------------
    sheet = book.active
    sheet.title = 'Settings'
    widths = {}
    for values in settings_rows(project_paths, metrics, prep, manifest,
                                category, source):
        sheet.append(list(values))
        for index, value in enumerate(values, start=1):
            widths[index] = max(widths.get(index, 0), len(str(value)))
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is not None:
                cell.font = Font(name=FONT)
    _style_header(sheet, 1, 4, openpyxl)
    # the block titles, which sit alone on their row
    for row in sheet.iter_rows(min_col=1, max_col=1):
        cell = row[0]
        if cell.value and sheet.cell(row=cell.row, column=2).value is None \
                and cell.row > 1:
            cell.font = Font(name=FONT, bold=True)
            cell.fill = PatternFill('solid', fgColor=BLOCK_FILL)
    _autofit(sheet, widths)
    sheet.freeze_panes = 'A2'

    # --- tab 2 ------------------------------------------------------------
    trials = [t for t in metrics.get('trials', []) if not t.get('excluded')]
    for trial in trials:
        trial['_projectID'] = metrics.get('projectID')
        trial['_analysisID'] = metrics.get('analysisID')
        trial['_tankID'] = metrics.get('tankID')
        trial['_category'] = category
    _write_table(book, 'Trials', trial_columns(), trials, openpyxl)

    # --- tab 3 ------------------------------------------------------------
    marks = metrics.get('marks') or {}
    new_bower = set(marks.get('newBower') or [])
    wall = set(marks.get('wallBuilding') or [])
    day_rows = []
    for trial in trials:
        events = trial.get('dayEvents') or {}
        for row in trial.get('dayVolumes') or []:
            entry = dict(row)
            entry['_projectID'] = metrics.get('projectID')
            entry['_analysisID'] = metrics.get('analysisID')
            entry['_category'] = category
            entry['_events'] = events.get(str(row['index'])) or {}
            light = _get(row, 'light', 'totalVolume', default=0) or 0
            dark = _get(row, 'dark', 'totalVolume', default=None)
            if dark is not None:
                entry['_total24'] = round(light + dark, 2)
                entry['_nightShare'] = round(dark / (light + dark), 4) \
                    if light + dark else None
            entry['_newBower'] = 'yes' if row['index'] in new_bower else ''
            entry['_wallBuilding'] = 'yes' if row['index'] in wall else ''
            day_rows.append(entry)
    _write_table(book, 'Days', day_columns(), day_rows, openpyxl)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    return path


def _write_table(book, title: str, columns, rows, openpyxl) -> None:
    from openpyxl.styles import Font
    sheet = book.create_sheet(title)
    sheet.append([name for name, _ in columns])
    widths = {i: len(name) for i, (name, _) in enumerate(columns, start=1)}
    for row in rows:
        values = []
        for _, getter in columns:
            try:
                value = getter(row)
            except Exception:
                value = None
            if isinstance(value, bool):
                value = 'yes' if value else 'no'
            values.append(value)
        sheet.append(values)
        for index, value in enumerate(values, start=1):
            if value is not None:
                widths[index] = max(widths.get(index, 0), len(str(value)))
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name=FONT)
    _style_header(sheet, 1, len(columns), openpyxl)
    _autofit(sheet, widths)
    sheet.freeze_panes = 'E2'
